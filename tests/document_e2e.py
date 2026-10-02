"""Real DeepL acceptance for the bundled document worker; never prints the API key."""
import json
import os
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

from docx import Document
from openpyxl import Workbook, load_workbook
from pptx import Presentation


ROOT = Path(__file__).resolve().parents[1]
WORKER = ROOT / "src-tauri/resources/scripts/translate_document.py"
PYTHON = ROOT / "src-tauri/resources/python/python.exe"
SOFFICE = ROOT / "src-tauri/resources/libreoffice/program/soffice.exe"


def load_key():
    for line in (ROOT / ".env").read_text(encoding="utf-8").splitlines():
        if line.startswith("DEEPL_API_KEY="):
            return line.partition("=")[2].strip().strip('"')
    raise RuntimeError("DEEPL_API_KEY is required for real document E2E acceptance.")


def translate(source, output, kind):
    env = os.environ | {"DEEPL_API_KEY": load_key()}
    subprocess.run([PYTHON, WORKER, source, output, kind, "ZH", "EN", json.dumps([])], check=True, env=env, capture_output=True)
    assert output.is_file() and output.stat().st_size > 0


def preview(path, directory):
    profile = directory / ("profile-" + path.suffix[1:])
    subprocess.run([SOFFICE, "--headless", f"-env:UserInstallation={profile.as_uri()}", "--convert-to", "pdf", "--outdir", directory, path], check=True, capture_output=True)
    assert path.with_suffix(".pdf").is_file()


def run():
    with tempfile.TemporaryDirectory(prefix="honsen-document-e2e-") as temporary:
        root = Path(temporary)
        source = root / "source.docx"; output = root / "zh_word.docx"
        document = Document(); document.add_heading("Electrical installation", 1); document.add_paragraph("Install the air handling unit.")
        document.sections[0].header.paragraphs[0].text = "Project specification"
        table = document.add_table(rows=1, cols=2); table.cell(0, 0).text = "Description"; table.cell(0, 1).text = "Quantity"
        document.save(source); translate(source, output, "word")
        result = Document(output); assert result.paragraphs[0].style.name.startswith("Heading") and "Electrical installation" not in "\n".join(p.text for p in result.paragraphs); preview(output, root)

        source = root / "source.pptx"; output = root / "zh_presentation.pptx"
        presentation = Presentation(); slide = presentation.slides.add_slide(presentation.slide_layouts[1]); slide.shapes.title.text = "Electrical installation"; slide.placeholders[1].text = "Install the air handling unit."
        presentation.save(source); translate(source, output, "presentation")
        assert "Electrical installation" not in "\n".join(shape.text for shape in Presentation(output).slides[0].shapes if hasattr(shape, "text")); preview(output, root)

        source = root / "source.xlsx"; output = root / "zh_spreadsheet.xlsx"
        workbook = Workbook(); sheet = workbook.active; sheet["A1"] = "Electrical installation"; sheet["B1"] = "Quantity"; sheet["B2"] = 2; sheet["B3"] = 3; sheet["B4"] = "=SUM(B2:B3)"; sheet.merge_cells("A5:B5"); sheet["A5"] = "Project specification"; workbook.save(source)
        translate(source, output, "spreadsheet")
        result = load_workbook(output, data_only=False); assert result.active["B4"].value == "=SUM(B2:B3)" and "A5:B5" in {str(item) for item in result.active.merged_cells.ranges} and result.active["A1"].value != "Electrical installation"; preview(output, root)

        source = root / "source.txt"; output = root / "zh_source.txt"; source.write_text("Electrical installation\nInstall the air handling unit.\n", encoding="utf-8")
        translate(source, output, "text"); assert "Electrical installation" not in output.read_text(encoding="utf-8")

        source = root / "source.md"; output = root / "zh_source.md"; source.write_text("# Electrical installation\n[Guide](https://example.com)\n`keep_code()`\n![Unit](unit.png)\n", encoding="utf-8")
        translate(source, output, "markdown"); text = output.read_text(encoding="utf-8"); assert "https://example.com" in text and "`keep_code()`" in text and "Electrical installation" not in text
        with zipfile.ZipFile(root / "zh_word.docx") as archive: assert "word/document.xml" in archive.namelist()
    print("Document E2E accepted: DOCX, PPTX, XLSX, TXT, Markdown.")


if __name__ == "__main__":
    run()
